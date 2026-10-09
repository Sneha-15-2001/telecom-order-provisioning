"""Generates Telecom-Order-Provisioning-Executive.pptx — high-level, no code. Run: python3 docs/generate-ppt-exec.py."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

RED = RGBColor(0xE4, 0x00, 0x00)
NAVY = RGBColor(0x0B, 0x1E, 0x4B)
DARK = RGBColor(0x20, 0x20, 0x20)
GRAY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF5, 0xF5, 0xF5)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def box(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def text(shape, value, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def cover(title, subtitle):
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY)
    box(s, Inches(0), Inches(0), Inches(13.333), Inches(0.12), RED)
    text(box(s, Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.6), NAVY),
         title, size=46, bold=True, color=WHITE)
    text(box(s, Inches(0.8), Inches(3.6), Inches(11.7), Inches(2.2), NAVY),
         subtitle, size=22, color=RGBColor(0xD7, 0xDE, 0xF0))
    return s


def points(kicker, title, bullets):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    box(s, Inches(0), Inches(0), Inches(13.333), Inches(0.12), RED)
    text(box(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.6), WHITE),
         kicker, size=16, bold=True, color=RED)
    text(box(s, Inches(0.8), Inches(1.1), Inches(11.7), Inches(1.0), WHITE),
         title, size=36, bold=True, color=NAVY)
    for i, b in enumerate(bullets):
        shape = box(s, Inches(0.8 + (i % 2) * 6.1), Inches(2.5 + (i // 2) * 1.35),
                    Inches(5.9), Inches(1.15), LIGHT)
        text(shape, b, size=17, color=DARK)
    return s


cover("One Connection,\nFully Automated",
      "How a customer order flows from signup to active service —\n"
      "validated, billed, provisioned and confirmed in one traceable journey.")

points("THE IDEA", "What the system does (in plain words)",
       ["A customer orders a mobile or broadband connection online.",
        "The system checks them, takes payment, reserves their number, switches it on, and texts them.",
        "If anything fails halfway, it undoes the half-done work automatically — no stuck customers.",
        "Every step is recorded, so any failure can be replayed and explained."])

points("HOW IT FEELS", "The journey in four moments",
       ["1 · Order placed — cart, eligibility and payment in one smooth flow.",
        "2 · Fulfilled — stock reserved, service activated, SMS delivered.",
        "3 · If trouble — order parks safely with the exact reason attached.",
        "4 · Always visible — customer, agent and engineer see the same timeline."])

points("THE AI ANGLE", "From incident ticket to answer in one click",
       ["Support pastes a ticket; the investigator traces it across systems by itself.",
        "It shows what broke, where, and the evidence — logs, records, timeline.",
        "It proposes the safe temporary fix and the permanent code fix for approval.",
        "Humans stay in charge: the AI recommends, people approve."])

points("VALUE", "Why this matters to the business",
       ["Faster resolution: minutes of tracing instead of hours across teams.",
        "Fewer repeat failures: every incident becomes a documented, replayable lesson.",
        "Safer changes: fixes are proposed with validation and rollback notes.",
        "Happier customers: no silent stuck orders, no duplicate messages, no bill shocks."])

points("ROADMAP", "Where it goes next",
       ["Now: full order journey + AI investigation on live simulated systems.",
        "Next: event-driven upgrades (Kafka) for scale and resilience.",
        "Then: one-command Docker deployment for demos anywhere.",
        "Vision: the same investigator pattern watching real production systems."])

cover("Thank You",
      "Telecom Order Provisioning System — built from scratch,\ndemoed end to end, investigated by AI.")

prs.save("/Users/sneha/telecom-order-provisioning/Telecom-Order-Provisioning-Executive.pptx")
print("saved")
